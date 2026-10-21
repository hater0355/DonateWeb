import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


source = Path(sys.argv[1])
target = Path(sys.argv[2])
data = json.loads(source.read_text(encoding="utf-8"))
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

for style_name, size in (("Title", 17), ("Heading 1", 13)):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(31, 61, 45)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(5)

doc.core_properties.title = data["title"]
doc.core_properties.subject = "Mẫu thỏa thuận điện tử giữa DonateWeb và Streamer"
doc.core_properties.author = "DonateWeb"
doc.core_properties.keywords = "DonateWeb, Streamer, hợp đồng điện tử"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(12)
p.paragraph_format.space_after = Pt(2)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Độc lập - Tự do - Hạnh phúc")
r.bold = True
r.font.size = Pt(12)
p.paragraph_format.space_after = Pt(10)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run(data["title"])
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(f"Phiên bản {data['version']}")
r.bold = True
r.font.size = Pt(10)
p.paragraph_format.space_after = Pt(10)

note = doc.add_table(rows=1, cols=1)
note.autofit = True
cell = note.cell(0, 0)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
shade = OxmlElement("w:shd")
shade.set(qn("w:fill"), "FFF4CE")
cell._tc.get_or_add_tcPr().append(shade)
cell.text = (
    "BẢN MẪU CHƯA SẴN SÀNG KÝ: cần điền tư cách pháp lý, địa chỉ và mã số thuế/mã định danh phù hợp của Bên A. "
    "Ứng dụng DonateWeb khóa thao tác chấp nhận cho đến khi chủ nền tảng hoàn tất các thông tin này. "
    "Bản Word này chưa phải bản hợp đồng đã ký của một Streamer cụ thể."
)
for paragraph in cell.paragraphs:
    paragraph.paragraph_format.space_after = Pt(2)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(120, 83, 0)

doc.add_paragraph("Các bên dự kiến giao kết hợp đồng điện tử theo các thông tin dưới đây:")

for heading, lines in [
    ("BÊN A - ĐƠN VỊ CUNG CẤP NỀN TẢNG", [
        "Tư cách pháp lý: [Bổ sung tư cách và căn cứ đăng ký kinh doanh]",
        "Họ và tên: Nguyễn Anh Hiếu",
        "Địa chỉ: [Bổ sung địa chỉ pháp lý]",
        "Mã số thuế/mã định danh/đăng ký: [Bổ sung mã số thật]",
        "Email liên hệ: cupitcookies@gmail.com",
        "Website: [Bổ sung địa chỉ website chính thức]",
    ]),
    ("BÊN B - STREAMER", [
        "Họ và tên theo CCCD: [Lấy từ hồ sơ định danh đã được duyệt]",
        "CCCD: [Số định danh của Streamer]",
        "Địa chỉ thường trú: [Lấy từ hồ sơ định danh]",
        "Email tài khoản: [Email đã đăng ký DonateWeb]",
        "Tên kênh và trang Streamer: [Tên kênh] / [Đường dẫn trang]",
    ]),
]:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(heading)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 61, 45)
    for line in lines:
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(2)

doc.add_paragraph(
    "Hai bên thống nhất các điều khoản sau. Nội dung, thời điểm và thông tin giao kết cụ thể của mỗi Streamer "
    "được ghi nhận trong bản hợp đồng điện tử do DonateWeb lưu sau khi Streamer chấp nhận."
)

for clause in data["sections"]:
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run(clause["title"])
    for text in clause["paragraphs"]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.widow_control = True

heading = doc.add_paragraph(style="Heading 1")
heading.add_run("Xác nhận giao kết điện tử")
doc.add_paragraph(
    "Bên B giao kết bằng tài khoản DonateWeb đã xác thực, xác nhận đã đọc và đồng ý toàn bộ nội dung của đúng phiên bản hợp đồng, "
    "nhập lại họ tên trùng với hồ sơ CCCD và xác nhận đã đọc Chính sách bảo mật và quyền riêng tư. Hệ thống lưu bản nội dung, "
    "mã SHA-256, thời điểm, tài khoản, địa chỉ IP và thông tin trình duyệt cần thiết để đối chiếu. Bản hợp đồng cụ thể được lưu "
    "trên DonateWeb và có thể được in hoặc lưu lại từ trang xác nhận sau khi ký."
)

doc.save(target)
print(target)
